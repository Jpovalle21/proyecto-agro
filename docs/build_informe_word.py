from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "INFORME_FRONTEND_AGRO_ANGELES.md"
OUTPUT = ROOT / "INFORME_FRONTEND_AGRO_ANGELES.docx"
ASSETS = ROOT / "informe-assets"

ACCENT = "003049"
ACCENT_2 = "00838F"
LIGHT_BLUE = "EAF6F8"
LIGHT_GRAY = "F4F6F8"
BORDER = "D9D9D9"
TEXT = RGBColor(15, 23, 42)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_fixed_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text_with_code(paragraph, text):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", 9.5)
            run.font.color.rgb = RGBColor(51, 65, 85)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, "Aptos", 10.5)


def format_paragraph(paragraph, before=0, after=7, line=1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_normal_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.style = "Body Text"
    add_text_with_code(paragraph, text)
    format_paragraph(paragraph)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    add_text_with_code(paragraph, text)
    format_paragraph(paragraph, after=4)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    add_text_with_code(paragraph, text)
    format_paragraph(paragraph, after=4)
    return paragraph


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_border(cell, "CBD5E1")
    set_cell_margins(cell, 150, 170, 150, 170)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(code.rstrip())
    set_run_font(run, "Consolas", 8.6)
    paragraph.paragraph_format.line_spacing = 1.05
    doc.add_paragraph()


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if r_index == 0:
                set_cell_shading(cell, ACCENT)
            elif r_index % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            set_run_font(run, "Aptos", 8.5 if len(value) > 80 else 9)
            if r_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = TEXT
    doc.add_paragraph()


def add_image(doc, alt, rel_path):
    image_path = (ROOT / rel_path).resolve()
    if not image_path.exists():
        add_normal_paragraph(doc, f"[Imagen no encontrada: {rel_path}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.35))
    format_paragraph(paragraph, before=4, after=3)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(alt)
    set_run_font(run, "Aptos", 9, bold=True)
    run.font.color.rgb = RGBColor(71, 85, 105)
    format_paragraph(caption, after=12)


def add_cover(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(22)
    run = paragraph.add_run("Informe del Frontend Agro Angeles")
    paragraph.style = "Title"
    set_run_font(run, "Aptos Display", 25, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Descripcion funcional y tecnica de los modulos desarrollados")
    set_run_font(run, "Aptos", 13, color="475569")
    format_paragraph(subtitle, after=26)

    meta_rows = [
        ["Fecha del informe", "7 de octubre de 2026"],
        ["Proyecto", "inventario-frontend"],
        ["Ubicacion", r"C:\Users\Usuario\OneDrive\Documentos\Frontend\inventario-frontend"],
        ["Tipo de documento", "Informe completo de desarrollo frontend"],
    ]
    table = doc.add_table(rows=len(meta_rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(meta_rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            set_cell_border(cell)
            set_cell_margins(cell, 170, 180, 170, 180)
            if col_index == 0:
                set_cell_shading(cell, ACCENT)
            else:
                set_cell_shading(cell, "F8FAFC" if row_index % 2 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(value)
            set_run_font(run, "Aptos", 9.5, bold=col_index == 0)
            run.font.color.rgb = RGBColor(255, 255, 255) if col_index == 0 else TEXT

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    intro_text = (
        "Este informe documenta el estado actual del frontend, los modulos creados, "
        "la razon de las decisiones tomadas, la forma en que funciona cada flujo y "
        "los requisitos pendientes para convertir el sistema en una herramienta productiva con backend."
    )
    add_text_with_code(intro, intro_text)
    format_paragraph(intro, before=18, after=18, line=1.15)

    add_image(doc, "Vista general del panel principal", "informe-assets/01-dashboard.png")
    doc.add_page_break()


def add_manual_toc(doc, headings):
    paragraph = doc.add_paragraph()
    paragraph.style = "Heading 1"
    paragraph.add_run("Contenido")
    format_paragraph(paragraph, after=10)
    for level, text in headings:
        if level != 2:
            continue
        item = doc.add_paragraph()
        item.paragraph_format.left_indent = Cm(0.2)
        item.paragraph_format.space_after = Pt(2)
        run = item.add_run(text)
        set_run_font(run, "Aptos", 9.7)
    doc.add_page_break()


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT

    for style_name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Heading 1" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    body = styles["Body Text"]
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    body.font.size = Pt(10.5)
    body.font.color.rgb = TEXT

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.2)
        style.font.color.rgb = TEXT


def add_footer(doc):
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Informe del Frontend Agro Angeles")
        set_run_font(run, "Aptos", 8.5, color="64748B")


def parse_table(lines, start):
    table_lines = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1
    parsed = []
    for line in table_lines:
        values = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", value or "") for value in values):
            continue
        parsed.append(values)
    return parsed, index


def extract_headings(lines):
    headings = []
    for line in lines:
        if line.startswith("## "):
            headings.append((2, line[3:].strip()))
    return headings


def build_docx():
    source_text = SOURCE.read_text(encoding="utf-8")
    lines = source_text.splitlines()
    doc = Document()
    setup_styles(doc)
    add_cover(doc)
    add_manual_toc(doc, extract_headings(lines))

    in_code = False
    code_lines = []
    skip_title_meta = True
    index = 0

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()

        if skip_title_meta:
            if line.startswith("## 1. "):
                skip_title_meta = False
            else:
                index += 1
                continue

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        image_match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", line)
        if image_match:
            add_image(doc, image_match.group(1), image_match.group(2))
            index += 1
            continue

        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        if line.startswith("### "):
            paragraph = doc.add_paragraph(line[4:].strip())
            paragraph.style = "Heading 3"
            index += 1
            continue

        if line.startswith("## "):
            paragraph = doc.add_paragraph(line[3:].strip())
            paragraph.style = "Heading 1"
            index += 1
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph(line[2:].strip())
            paragraph.style = "Heading 1"
            index += 1
            continue

        bullet_match = re.match(r"[-*]\s+(.*)", line)
        if bullet_match:
            add_bullet(doc, bullet_match.group(1).strip())
            index += 1
            continue

        numbered_match = re.match(r"\d+\.\s+(.*)", line)
        if numbered_match:
            add_numbered(doc, numbered_match.group(1).strip())
            index += 1
            continue

        add_normal_paragraph(doc, line)
        index += 1

    add_footer(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT)
