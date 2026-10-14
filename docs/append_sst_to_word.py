from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


DOCS = Path(__file__).resolve().parent
INPUT = DOCS / "INFORME_FRONTEND_AGRO_ANGELES_PROYECTOS_ADICIONALES_SST_BASE.docx"
OUTPUT = DOCS / "INFORME_FRONTEND_AGRO_ANGELES_COMPLETO_CON_SST.docx"
ASSETS = DOCS / "informe-assets"

ACCENT = "003049"
SST = "E67E22"
BORDER = "D9D9D9"
TEXT = RGBColor(15, 23, 42)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_paragraph(paragraph, before=0, after=7, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text)
    paragraph.style = f"Heading {level}"
    format_paragraph(paragraph, before=10 if level == 1 else 7, after=6)
    return paragraph


def add_paragraph(doc, text, style="Body Text", before=0, after=7):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, "Aptos", 10.5)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, before=before, after=after)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, "Aptos", 10.2)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, after=4)
    return paragraph


def add_table(doc, headers, rows, header_fill=ACCENT):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell, 140, 140, 140, 140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, "Aptos", 8.8, bold=True, color="FFFFFF")

    for row_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            cell = table.cell(row_index, col)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            run = paragraph.add_run(value)
            set_run_font(run, "Aptos", 8.45)
            run.font.color.rgb = TEXT
    doc.add_paragraph()
    return table


def add_image(doc, image_path, caption, width=6.25):
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    format_paragraph(paragraph, before=3, after=3)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    set_run_font(run, "Aptos", 9, bold=True)
    run.font.color.rgb = RGBColor(71, 85, 105)
    format_paragraph(caption_paragraph, after=10)


def insert_toc_entry(doc):
    matches = [p for p in doc.paragraphs if p.text.strip() == "24. Proyecto independiente Pagina Web Institucional Agro Angeles"]
    if not matches:
        return
    toc_match = matches[0]
    new_p = OxmlElement("w:p")
    toc_match._p.addnext(new_p)
    paragraph = Paragraph(new_p, toc_match._parent)
    run = paragraph.add_run("25. Actualizacion del frontend Modulo SST")
    set_run_font(run, "Aptos", 9.7)
    paragraph.paragraph_format.left_indent = toc_match.paragraph_format.left_indent
    paragraph.paragraph_format.space_after = Pt(2)


def append_sst_section(doc):
    doc.add_page_break()
    add_heading(doc, "25. Actualizacion del frontend Modulo SST", 1)
    add_paragraph(
        doc,
        "Esta seccion documenta la incorporacion del modulo de Seguridad y Salud en el Trabajo al frontend. "
        "El modulo se diseno a partir de los archivos de Excel que actualmente usa el area para controlar incapacidades y examenes medicos ocupacionales. "
        "Ademas, se agrego un canal para que otras areas puedan enviar solicitudes o reportar incidentes y que todo quede registrado dentro del sistema.",
    )

    add_heading(doc, "25.1 Para que se hizo", 2)
    add_paragraph(
        doc,
        "El modulo SST se hizo para que la coordinacion de seguridad y salud pueda consultar informacion critica sin depender de varios archivos sueltos. "
        "La idea es que incapacidades, examenes medicos, vencimientos, pendientes y reportes de areas queden visibles en una misma interfaz, con busqueda, filtros y estados.",
    )
    add_paragraph(
        doc,
        "Tambien se planteo como un puente entre SST y las demas areas. Produccion, mantenimiento, calidad, almacen u otras areas pueden reportar incidentes, actos inseguros, condiciones inseguras o solicitudes de apoyo, dejando evidencia escrita para seguimiento y cierre.",
    )

    add_heading(doc, "25.2 Que se desarrollo", 2)
    for item in [
        "Dashboard SST con indicadores de personal activo, solicitudes abiertas, pagos pendientes, examenes vencidos, por vencer, vigentes, faltantes y sin fecha.",
        "Pantalla de incapacidades con registro EPS o ARL, persona, cedula, tipo, dias, fecha de inicio, entidad, valor, cobro, pago y observaciones.",
        "Pantalla de examenes medicos ocupacionales con estados vencido, por vencer, vigente, faltante y sin fecha.",
        "Formulario para actualizar examenes, calculando automaticamente la fecha de vencimiento a un ano desde la fecha del examen.",
        "Pantalla de solicitudes e incidentes para reportes de otras areas, con tipo, area, responsable, prioridad, descripcion, estado y acciones de SST.",
        "Sidebar SST integrado al mismo estilo de navegacion que ya tienen los modulos existentes.",
        "Persistencia funcional con localStorage para probar el flujo completo mientras se conecta backend.",
        "Datos iniciales basados en los dos Excel recibidos para que el prototipo represente informacion real del area.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "25.3 Como funciona", 2)
    add_paragraph(
        doc,
        "El modulo se agrego dentro del frontend React existente. La ruta principal es /app/sst/dashboard y desde el sidebar se puede entrar a incapacidades, examenes medicos y solicitudes. "
        "Cada pantalla carga datos iniciales desde un archivo de almacenamiento local del modulo y guarda los cambios en localStorage. Esto permite probar altas, busquedas, filtros, cambios de estado y registro de acciones sin depender todavia de una API.",
    )
    add_table(
        doc,
        ["Vista", "Funcion", "Datos principales"],
        [
            ["Dashboard SST", "Resume prioridades del area y muestra accesos rapidos.", "Personal activo, solicitudes abiertas, pagos pendientes, vencidos, por vencer, vigentes, faltantes y sin fecha."],
            ["Incapacidades", "Permite registrar y consultar incapacidades EPS y ARL.", "Persona, cedula, tipo, dias, fecha, entidad, valor, si se cobro, si fue pagada y observaciones."],
            ["Examenes medicos", "Controla vigencias y pendientes de examenes ocupacionales.", "Nombre, cedula, cargo, tipo de examen, fecha, vencimiento, dias, concepto y observacion."],
            ["Solicitudes e incidentes", "Recibe reportes de otras areas y permite seguimiento SST.", "Tipo, area, quien reporta, prioridad, titulo, descripcion, estado y acciones realizadas."],
        ],
        header_fill=SST,
    )

    add_heading(doc, "25.4 Relacion con los Excel actuales", 2)
    add_table(
        doc,
        ["Archivo usado por SST", "Informacion que contiene", "Como se llevo al frontend"],
        [
            ["Informe de incapacidades 2026", "Hojas separadas para ARL y EPS con persona, cedula, tipo de incapacidad, dias, fecha, cobro, entidad, valor, pago y observaciones.", "Se convirtio en una pantalla de incapacidades con registro, filtros por EPS/ARL, valores acumulados y estado de pago."],
            ["Seguimiento examenes medicos activos", "Resumen por personal activo, faltantes, sin fecha, vencidos, por vencer, vigentes y no activos.", "Se convirtio en dashboard de estados y pantalla de examenes con busqueda, filtros, vencimientos y actualizacion de examenes."],
            ["Necesidad nueva de reportes de areas", "No estaba centralizada en los Excel, pero es clave para dejar evidencia escrita.", "Se agrego una vista de solicitudes e incidentes para que las areas reporten casos y SST registre acciones y cierre."],
        ],
        header_fill=SST,
    )

    add_heading(doc, "25.5 Por que se hizo de esa manera", 2)
    add_table(
        doc,
        ["Decision", "Razon"],
        [
            ["Usar los Excel como base de campos", "Evita inventar informacion y mantiene el flujo cercano a lo que SST ya maneja."],
            ["Separar incapacidades y examenes", "Son procesos diferentes: incapacidades requieren cobro y pago, mientras examenes requieren vigencia y alertas de vencimiento."],
            ["Agregar dashboard primero", "SST necesita ver prioridades rapidamente antes de entrar al detalle."],
            ["Incluir solicitudes e incidentes", "Permite que otras areas reporten casos sin depender de mensajes sueltos o conversaciones informales."],
            ["Guardar en localStorage por ahora", "Permite probar el modulo completo en frontend mientras se define backend, usuarios reales y base de datos."],
            ["Mantener el mismo sidebar", "Conserva coherencia visual con Almacen, Calidad y los demas modulos del sistema."],
        ],
        header_fill=SST,
    )

    add_heading(doc, "25.6 Evidencia visual", 2)
    add_image(doc, ASSETS / "sst-dashboard-desktop.png", "Dashboard del modulo SST con indicadores y prioridades", width=6.35)
    add_image(doc, ASSETS / "sst-incapacidades-desktop.png", "Pantalla de incapacidades con referencia al Excel y formulario de registro", width=6.35)
    add_image(doc, ASSETS / "sst-examenes-desktop.png", "Pantalla de examenes medicos ocupacionales con estados de seguimiento", width=6.35)
    add_image(doc, ASSETS / "sst-solicitudes-desktop.png", "Pantalla de solicitudes e incidentes para reportes desde otras areas", width=6.35)

    add_heading(doc, "25.7 Resultado para el proyecto", 2)
    add_paragraph(
        doc,
        "Con este modulo, el frontend avanza hacia una plataforma mas completa para las areas administrativas y operativas. SST ya no queda representado solo como un acceso vacio, sino como un flujo funcional con datos, pantallas, formularios, filtros, estados y registro de acciones.",
    )
    add_paragraph(
        doc,
        "El siguiente paso recomendado es conectar el modulo SST a backend para que las incapacidades, examenes y reportes se guarden en base de datos, tengan permisos por area, adjuntos, trazabilidad por usuario y exportaciones oficiales.",
    )


def main():
    doc = Document(INPUT)
    insert_toc_entry(doc)
    append_sst_section(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
