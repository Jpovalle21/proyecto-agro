from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


DOCS = Path(__file__).resolve().parent
INPUT = DOCS / "INFORME_FRONTEND_AGRO_ANGELES.docx"
OUTPUT = DOCS / "INFORME_FRONTEND_AGRO_ANGELES_CON_CRONOGRAMA.docx"
CRONOGRAMA_ROOT = Path(r"C:\Users\Usuario\OneDrive\Documentos\Cronograma")

ACCENT = "003049"
BORDER = "D9D9D9"
TEXT = RGBColor(15, 23, 42)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_paragraph(paragraph, before=0, after=7, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_paragraph(doc, text, style="Body Text", before=0, after=7):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, "Aptos", 10.5)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, before=before, after=after)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text)
    paragraph.style = f"Heading {level}"
    format_paragraph(paragraph, before=10 if level == 1 else 7, after=6)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, "Aptos", 10.2)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, after=4)
    return paragraph


def add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        set_cell_shading(cell, ACCENT)
        set_cell_border(cell)
        set_cell_margins(cell, 140, 140, 140, 140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        set_run_font(run, "Aptos", 8.8, bold=True, color="FFFFFF")

    for row_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            cell = table.cell(row_index, col)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            run = paragraph.add_run(value)
            set_run_font(run, "Aptos", 8.5)
            run.font.color.rgb = TEXT
    doc.add_paragraph()
    return table


def add_image(doc, image_path, caption, width=6.35):
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    format_paragraph(paragraph, before=3, after=3)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    set_run_font(run, "Aptos", 9, bold=True)
    run.font.color.rgb = RGBColor(71, 85, 105)
    format_paragraph(caption_paragraph, after=10)


def insert_toc_entry(doc):
    matches = [p for p in doc.paragraphs if p.text.strip() == "22. Conclusiones"]
    if not matches:
        return
    toc_match = matches[0]
    new_p = OxmlElement("w:p")
    toc_match._p.addnext(new_p)
    paragraph = Paragraph(new_p, toc_match._parent)
    run = paragraph.add_run("23. Anexo Proyecto Cronograma Semanal Agro Angeles")
    set_run_font(run, "Aptos", 9.7)
    paragraph.paragraph_format.left_indent = toc_match.paragraph_format.left_indent
    paragraph.paragraph_format.space_after = Pt(2)


def append_cronograma_section(doc):
    doc.add_page_break()
    add_heading(doc, "23. Anexo Proyecto Cronograma Semanal Agro Angeles", 1)
    add_paragraph(
        doc,
        "Este anexo documenta el trabajo realizado en el sistema web de cronograma semanal. "
        "Se incluye para mostrar con claridad que el desarrollo no se limito a una interfaz visual, "
        "sino que incorporo frontend, backend, base de datos, roles, reportes, recordatorios por correo "
        "y preparacion de despliegue.",
    )

    add_heading(doc, "23.1 Para que se hizo", 2)
    add_paragraph(
        doc,
        "El cronograma se hizo para centralizar la planificacion, registro, consulta y control de actividades por area. "
        "Su finalidad es que cada jefe o auxiliar pueda registrar lo que debe realizar, mientras Auditoria Externa "
        "tiene una vista general del cumplimiento semanal y puede revisar actividades por area, responsable, estado e historico.",
    )
    add_paragraph(
        doc,
        "Antes del sistema, el seguimiento dependia de archivos separados, avisos manuales y comunicacion directa entre areas. "
        "Eso hacia mas dificil saber que actividades estaban pendientes, cuales se habian cumplido y por que algunas no se realizaron. "
        "Con el aplicativo, la informacion queda centralizada y consultable desde varios computadores.",
    )

    add_heading(doc, "23.2 Que se desarrollo", 2)
    for item in [
        "Login con correo y contrasena, validado por backend y protegido con JWT.",
        "Panel operativo con resumen de actividades, pendientes, actividades del dia, pendientes de ayer, frase del dia y accesos rapidos.",
        "Cronograma semanal en formato de 24 horas para visualizar actividades por fecha, area y estado.",
        "Gestion de actividades con estados pendiente, cumplida y no cumplida, incluyendo motivo de incumplimiento.",
        "Actividades favoritas para reutilizar tareas repetitivas sin volver a escribir toda la descripcion.",
        "Recordatorios personales unicos o recurrentes semanalmente.",
        "Modulo de capacitaciones con areas convocadas, notificacion por correo, generacion de actividades asociadas y cancelacion con motivo.",
        "Comentarios administrativos sobre actividades finalizadas para que Auditoria pueda dejar observaciones.",
        "Vista de incumplimientos para revisar responsables, areas, fechas y motivos.",
        "Historicos con archivado automatico para que el cronograma activo no crezca indefinidamente.",
        "Exportaciones en Excel y PDF para cronograma general, cronograma por area y reportes historicos.",
        "Gestion tecnica de usuarios y areas desde el usuario de Sistemas con permisos administrativos.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "23.3 Como funciona", 2)
    add_paragraph(
        doc,
        "El flujo general es: el usuario entra desde el navegador, el frontend React muestra las pantallas del sistema, "
        "las acciones se envian a una API Express, el backend valida permisos y guarda la informacion mediante Prisma en MySQL. "
        "Cuando una actividad, capacitacion o recordatorio requiere notificacion, el backend envia correos automaticos por SMTP.",
    )
    add_table(
        doc,
        ["Capa", "Tecnologia", "Funcion dentro del sistema"],
        [
            ["Frontend", "React, Vite, React Router, React Query, Recharts", "Interfaz de login, paneles, cronograma, recordatorios, capacitaciones, usuarios, reportes y exportaciones."],
            ["Backend", "Node.js, Express, TypeScript", "API REST que valida solicitudes, aplica reglas de negocio y protege rutas por rol."],
            ["Base de datos", "MySQL con Prisma", "Guarda areas, usuarios, actividades, favoritos, recordatorios, capacitaciones, historicos y auditoria."],
            ["Seguridad", "JWT, bcrypt, helmet, validaciones", "Controla sesiones, contrasenas, acceso por rol y endurecimiento basico de la API."],
            ["Correos", "Nodemailer con Gmail SMTP", "Envia recuperacion de contrasena, recordatorios, capacitaciones y cancelaciones."],
            ["Despliegue", "Hostinger Cloud, variables de entorno y scripts npm", "Permite publicar el sistema con base de datos, migraciones y compilacion preparada."],
        ],
    )

    add_heading(doc, "23.4 Por que se hizo de esa manera", 2)
    add_table(
        doc,
        ["Decision", "Razon"],
        [
            ["Monorepo con apps api y web", "Permite mantener frontend y backend en un solo proyecto, con comandos compartidos de desarrollo, build, despliegue y mantenimiento."],
            ["Backend real con Express y Prisma", "El cronograma necesita datos compartidos por varios usuarios; por eso no bastaba guardar informacion solo en el navegador."],
            ["MySQL como base de datos", "Da persistencia centralizada para actividades, usuarios, areas, capacitaciones, historicos y auditoria."],
            ["Roles diferenciados", "Auditoria, Sistemas y usuarios operativos necesitan permisos distintos. Separar roles evita accesos indebidos y mantiene claridad operativa."],
            ["JWT para sesion segura", "Permite proteger rutas y mantener la sesion del usuario durante el uso del sistema."],
            ["Recordatorios y correos automaticos", "Reducen la dependencia de avisos manuales y ayudan a que las areas recuerden actividades, capacitaciones y cambios."],
            ["Historicos y limpieza automatica", "Evitan que el cronograma activo se vuelva pesado, pero conservan trazabilidad durante el periodo configurado."],
            ["Exportaciones PDF y Excel", "Auditoria necesita descargar evidencia, revisar cronogramas por area y compartir reportes fuera del sistema."],
        ],
    )

    add_heading(doc, "23.5 Roles implementados", 2)
    add_table(
        doc,
        ["Rol", "Responsabilidad principal"],
        [
            ["Administrador auditor", "Revisa dashboard general, cronogramas por area, incumplimientos, historicos, exportaciones y comentarios administrativos."],
            ["Usuario de Sistemas con permisos tecnicos", "Administra usuarios, areas, estados de cuenta y contrasenas temporales sin duplicar perfiles."],
            ["Usuario operativo", "Crea actividades, recordatorios, capacitaciones, favoritos y finaliza actividades como cumplidas o no cumplidas."],
        ],
    )

    add_heading(doc, "23.6 Evidencia tecnica del trabajo realizado", 2)
    add_table(
        doc,
        ["Elemento", "Evidencia dentro del proyecto"],
        [
            ["Aplicacion web", "apps/web con React, Vite, rutas, componentes, dashboard, cronograma, capacitaciones, recordatorios y reportes."],
            ["API", "apps/api con Express, rutas REST, middleware de autenticacion, servicios de correo, auditoria, archivado y recordatorios."],
            ["Base de datos", "Prisma schema con tablas Area, User, Activity, ActivityFavorite, PersonalReminder, Training, TrainingArea, ActivityHistory y AuditLog."],
            ["Despliegue", "Scripts deploy:prepare, deploy:init, build, start y documentacion para Hostinger Cloud."],
            ["Mantenimiento", "Documentos de backups, MySQL, SMTP Gmail, persistencia de datos, checklist de puesta en marcha y guias de operacion."],
        ],
    )

    add_heading(doc, "23.7 Capturas del sistema", 2)
    add_image(
        doc,
        CRONOGRAMA_ROOT / "desktop-dashboard.png",
        "Dashboard del sistema de cronograma semanal",
        width=6.35,
    )
    add_image(
        doc,
        CRONOGRAMA_ROOT / "mobile-schedule.png",
        "Vista movil del cronograma semanal",
        width=3.1,
    )

    add_heading(doc, "23.8 Resultado para la empresa", 2)
    add_paragraph(
        doc,
        "El resultado es una herramienta empresarial preparada para que las areas registren sus actividades y Auditoria pueda consultar cumplimiento, "
        "incumplimientos e historicos sin depender de archivos separados. El sistema tambien deja bases tecnicas importantes: autenticacion, roles, "
        "base de datos, correos, reportes, exportaciones, scripts de despliegue y documentacion de mantenimiento.",
    )
    add_paragraph(
        doc,
        "Con este desarrollo se demuestra capacidad para construir un sistema completo de uso interno, no solo una vista. El proyecto integra experiencia de usuario, "
        "organizacion por roles, persistencia de datos, automatizaciones y preparacion para despliegue real.",
    )


def main():
    doc = Document(INPUT)
    insert_toc_entry(doc)
    append_cronograma_section(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
