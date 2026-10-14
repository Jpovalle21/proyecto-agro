from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageOps


DOCS = Path(__file__).resolve().parent
INPUT = DOCS / "INFORME_FRONTEND_AGRO_ANGELES.docx"
OUTPUT = DOCS / "INFORME_FRONTEND_AGRO_ANGELES_PROYECTOS_ADICIONALES.docx"
ASSETS = DOCS / "informe-assets"

CRONOGRAMA_ROOT = Path(r"C:\Users\Usuario\OneDrive\Documentos\Cronograma")
WEB_ROOT = Path(r"C:\Users\Usuario\OneDrive\Documentos\Pagina web\pagina-web")
WEB_IMG = WEB_ROOT / "img"

ACCENT = "003049"
ACCENT_2 = "780000"
BORDER = "D9D9D9"
TEXT = RGBColor(15, 23, 42)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_paragraph(paragraph, before=0, after=7, line=1.08):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_paragraph(doc, text, style="Body Text", before=0, after=7):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, "Aptos", 10.5)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, before=before, after=after)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text)
    paragraph.style = f"Heading {level}"
    format_paragraph(paragraph, before=10 if level == 1 else 7, after=6)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    set_run_font(run, "Aptos", 10.2)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, after=4)
    return paragraph


def add_table(doc, headers, rows, header_fill=ACCENT):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    set_repeat_table_header(table.rows[0])

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell, 140, 140, 140, 140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, "Aptos", 8.8, bold=True, color="FFFFFF")

    for row_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            cell = table.cell(row_index, col)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            run = paragraph.add_run(value)
            set_run_font(run, "Aptos", 8.45)
            run.font.color.rgb = TEXT
    doc.add_paragraph()
    return table


def add_image(doc, image_path, caption, width=6.25):
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    format_paragraph(paragraph, before=3, after=3)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    set_run_font(run, "Aptos", 9, bold=True)
    run.font.color.rgb = RGBColor(71, 85, 105)
    format_paragraph(caption_paragraph, after=10)


def create_web_collage():
    ASSETS.mkdir(parents=True, exist_ok=True)
    output = ASSETS / "pagina-web-evidencia-visual.jpg"
    selections = [
        ("Inicio", WEB_IMG / "Imagen 1.jpeg"),
        ("Nuestra empresa", WEB_IMG / "Frame 3.png"),
        ("Beneficio", WEB_IMG / "Imagen 7.jpeg"),
        ("Trabaja", WEB_IMG / "Frame 5.png"),
        ("Contacto", WEB_IMG / "Imagen 8.jpeg"),
        ("Identidad", WEB_IMG / "Logo 1.png"),
    ]
    cell_w, cell_h = 560, 330
    label_h = 44
    gap = 18
    sheet_w = cell_w * 2 + gap * 3
    sheet_h = (cell_h + label_h) * 3 + gap * 4
    sheet = Image.new("RGB", (sheet_w, sheet_h), "white")
    draw = ImageDraw.Draw(sheet)

    for index, (label, path) in enumerate(selections):
        if not path.exists():
            continue
        img = Image.open(path)
        img = ImageOps.exif_transpose(img).convert("RGB")
        img.thumbnail((cell_w, cell_h), Image.LANCZOS)
        col = index % 2
        row = index // 2
        x = gap + col * (cell_w + gap)
        y = gap + row * (cell_h + label_h + gap)
        draw.rectangle((x, y, x + cell_w, y + cell_h + label_h), outline=(217, 217, 217), width=2)
        bg_x = x + (cell_w - img.width) // 2
        bg_y = y + (cell_h - img.height) // 2
        sheet.paste(img, (bg_x, bg_y))
        draw.rectangle((x, y + cell_h, x + cell_w, y + cell_h + label_h), fill=(248, 250, 252))
        draw.text((x + 16, y + cell_h + 13), label, fill=(15, 23, 42))

    sheet.save(output, quality=88, optimize=True)
    return output


def insert_toc_entries(doc):
    matches = [p for p in doc.paragraphs if p.text.strip() == "22. Conclusiones"]
    if not matches:
        return
    toc_match = matches[0]
    entries = [
        "PARTE II Proyectos adicionales desarrollados",
        "23. Proyecto independiente Cronograma Semanal Agro Angeles",
        "24. Proyecto independiente Pagina Web Institucional Agro Angeles",
    ]
    anchor = toc_match._p
    for entry in entries:
        new_p = OxmlElement("w:p")
        anchor.addnext(new_p)
        paragraph = Paragraph(new_p, toc_match._parent)
        run = paragraph.add_run(entry)
        set_run_font(run, "Aptos", 9.7, bold=entry.startswith("PARTE II"))
        paragraph.paragraph_format.left_indent = toc_match.paragraph_format.left_indent
        paragraph.paragraph_format.space_after = Pt(2)
        anchor = new_p


def append_part_divider(doc):
    doc.add_page_break()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PARTE II")
    set_run_font(run, "Aptos", 18, bold=True, color=ACCENT)
    format_paragraph(paragraph, before=60, after=4)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Proyectos adicionales desarrollados")
    set_run_font(run, "Aptos", 24, bold=True)
    run.font.color.rgb = TEXT
    format_paragraph(paragraph, after=12)

    add_paragraph(
        doc,
        "Esta parte separa claramente dos desarrollos adicionales del informe principal: "
        "el sistema de Cronograma Semanal y la Pagina Web Institucional. Ambos se documentan como proyectos independientes, "
        "con su objetivo, forma de funcionamiento, decisiones tomadas y evidencia tecnica.",
    )
    add_table(
        doc,
        ["Proyecto", "Para que se hizo", "Resultado principal"],
        [
            ["Cronograma Semanal", "Controlar actividades por area, cumplimiento, incumplimientos, recordatorios y reportes.", "Sistema web con frontend, backend, base de datos, roles, correos, historicos y exportaciones."],
            ["Pagina Web Institucional", "Presentar Agro Angeles hacia clientes, candidatos y visitantes externos.", "Sitio web responsive con paginas institucionales, contacto, chatbot WhatsApp y postulacion laboral funcional."],
        ],
        header_fill=ACCENT_2,
    )


def append_cronograma_section(doc):
    doc.add_page_break()
    add_heading(doc, "23. Proyecto independiente Cronograma Semanal Agro Angeles", 1)
    add_paragraph(
        doc,
        "Este proyecto se desarrollo aparte del frontend interno de modulos empresariales. Su proposito es controlar el cronograma semanal de actividades, "
        "registrar cumplimiento por area y entregar a Auditoria Externa una herramienta para revisar, comentar y descargar informacion operativa.",
    )

    add_heading(doc, "23.1 Para que se hizo", 2)
    add_paragraph(
        doc,
        "El cronograma se hizo para centralizar la planificacion, registro, consulta y control de actividades por area. "
        "Cada jefe o auxiliar puede registrar lo que debe realizar, mientras Auditoria Externa puede revisar el cumplimiento general, filtrar por area, "
        "consultar historicos y analizar actividades incumplidas.",
    )
    add_paragraph(
        doc,
        "La necesidad principal era dejar de depender de archivos separados, avisos manuales y seguimiento por comunicacion directa. "
        "El sistema permite que la informacion quede organizada en una base de datos y pueda consultarse desde diferentes equipos.",
    )

    add_heading(doc, "23.2 Que se desarrollo", 2)
    for item in [
        "Login con correo y contrasena, validado por backend y protegido con JWT.",
        "Dashboard para usuarios y auditoria con resumen de actividades, pendientes y cumplimiento.",
        "Cronograma semanal en formato de 24 horas para visualizar actividades por fecha, area y estado.",
        "Gestion de actividades con estados pendiente, cumplida y no cumplida, incluyendo motivo de incumplimiento.",
        "Favoritos para reutilizar actividades repetitivas y evitar escritura manual innecesaria.",
        "Recordatorios personales unicos o recurrentes semanalmente.",
        "Modulo de capacitaciones con areas convocadas, correos y actividades asociadas.",
        "Comentarios administrativos sobre actividades finalizadas.",
        "Vistas de incumplimientos, historicos y exportaciones en Excel y PDF.",
        "Gestion tecnica de usuarios y areas desde el usuario de Sistemas con permisos administrativos.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "23.3 Como funciona", 2)
    add_paragraph(
        doc,
        "El usuario entra desde el navegador y utiliza una interfaz React. Las acciones se envian a una API Express, el backend valida permisos por rol "
        "y guarda la informacion mediante Prisma en MySQL. Cuando se requiere notificar una actividad, una capacitacion o un recordatorio, el backend envia correos automaticos por SMTP.",
    )
    add_table(
        doc,
        ["Capa", "Tecnologia", "Funcion dentro del sistema"],
        [
            ["Frontend", "React, Vite, React Router, React Query, Recharts", "Muestra login, paneles, cronograma, recordatorios, capacitaciones, usuarios, reportes y exportaciones."],
            ["Backend", "Node.js, Express, TypeScript", "Expone la API REST, valida solicitudes y aplica reglas del negocio."],
            ["Base de datos", "MySQL con Prisma", "Guarda areas, usuarios, actividades, favoritos, recordatorios, capacitaciones, historicos y auditoria."],
            ["Seguridad", "JWT, bcrypt, helmet y validaciones", "Protege sesiones, contrasenas y rutas por rol."],
            ["Correos", "Nodemailer con Gmail SMTP", "Envia recuperacion de contrasena, recordatorios, capacitaciones y cancelaciones."],
            ["Despliegue", "Hostinger Cloud y scripts npm", "Prepara build, migraciones, variables de entorno y ejecucion en servidor."],
        ],
    )

    add_heading(doc, "23.4 Por que se hizo de esa manera", 2)
    add_table(
        doc,
        ["Decision", "Razon"],
        [
            ["Separar frontend y backend dentro de un monorepo", "Permite mantener interfaz y API en un mismo proyecto, pero con responsabilidades claras."],
            ["Usar base de datos MySQL", "El cronograma requiere informacion compartida por multiples usuarios, historicos y trazabilidad."],
            ["Manejar roles", "Auditoria, Sistemas y usuarios operativos necesitan permisos diferentes."],
            ["Usar correos automaticos", "Reduce el seguimiento manual y mejora la recordacion de actividades y capacitaciones."],
            ["Incluir historicos y exportaciones", "Auditoria necesita consultar evidencia y descargar reportes fuera del sistema."],
        ],
    )

    add_heading(doc, "23.5 Evidencia del desarrollo", 2)
    add_table(
        doc,
        ["Elemento", "Evidencia"],
        [
            ["Aplicacion web", "apps/web con React, rutas, componentes, dashboard, cronograma y reportes."],
            ["API", "apps/api con rutas REST, autenticacion, servicios de correo, auditoria y archivado."],
            ["Base de datos", "Prisma schema con modelos para usuarios, areas, actividades, recordatorios, capacitaciones e historicos."],
            ["Documentacion", "Carpeta docs con guia completa del proyecto, despliegue, base de datos, SMTP y mantenimiento."],
        ],
    )

    add_heading(doc, "23.6 Capturas del sistema", 2)
    add_image(doc, CRONOGRAMA_ROOT / "desktop-dashboard.png", "Dashboard del sistema de cronograma semanal", width=6.35)
    add_image(doc, CRONOGRAMA_ROOT / "mobile-schedule.png", "Vista movil del cronograma semanal", width=3.0)

    add_heading(doc, "23.7 Resultado", 2)
    add_paragraph(
        doc,
        "El resultado fue un sistema completo de uso interno para ordenar actividades, reducir perdida de informacion y facilitar revision por auditoria. "
        "No quedo planteado solo como diseno visual: incluye autenticacion, permisos, datos persistentes, correos, historicos, reportes y despliegue preparado.",
    )


def append_web_section(doc):
    collage = create_web_collage()
    doc.add_page_break()
    add_heading(doc, "24. Proyecto independiente Pagina Web Institucional Agro Angeles", 1)
    add_paragraph(
        doc,
        "La pagina web se desarrollo como un proyecto independiente orientado a la presencia publica de Agro Angeles. "
        "A diferencia de los sistemas internos, esta web esta pensada para clientes, visitantes y personas interesadas en trabajar con la empresa.",
    )

    add_heading(doc, "24.1 Para que se hizo", 2)
    add_paragraph(
        doc,
        "La pagina web se hizo para que Agro Angeles tuviera una presentacion digital clara, profesional y facil de consultar. "
        "El sitio comunica la identidad de la empresa, su historia, sus servicios, sus procesos de beneficio avicola, sus canales de contacto y la opcion de postulacion laboral.",
    )
    add_paragraph(
        doc,
        "Tambien se penso como una herramienta practica: no solo muestra informacion, sino que permite recibir mensajes, dirigir visitantes a WhatsApp y captar hojas de vida mediante un formulario funcional.",
    )

    add_heading(doc, "24.2 Que se desarrollo", 2)
    for item in [
        "Pagina de inicio con hero institucional, propuesta de valor, seccion quienes somos, fortalezas, preguntas frecuentes y llamado a contacto.",
        "Pagina Nuestra Empresa con historia, mision, vision, linea de tiempo y contenido corporativo.",
        "Pagina Beneficio con proceso avicola, tarjetas interactivas, estadisticas y seccion de seguridad alimentaria.",
        "Pagina Trabaja con Nosotros con modal de postulacion, campos condicionales, autorizacion de datos y carga de hoja de vida.",
        "Pagina Contacto con informacion de ubicacion, formulario de contacto y mapa embebido de Google Maps.",
        "Navbar compartido, menu responsive tipo hamburguesa y boton de contacto.",
        "Chatbot flotante conectado a WhatsApp con opciones para hojas de vida y servicios de maquila, congelacion y transporte.",
        "Formulario laboral con calculo automatico de puntaje, clasificacion del candidato, carga de archivo a Cloudinary y envio mediante EmailJS.",
        "Formulario de contacto conectado a FormSubmit para enviar mensajes al correo de gestion humana.",
        "Archivo ZIP preparado para subir el sitio a Hostinger.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "24.3 Como funciona", 2)
    add_paragraph(
        doc,
        "El sitio funciona como una web estatica construida con HTML, CSS y JavaScript. Cada pagina tiene su archivo HTML, comparte estilos desde style.css, "
        "usa imagenes locales de la carpeta img y reutiliza el chatbot desde js/chatbot.js y css/chatbot.css. Las funciones que requieren comunicacion externa se conectan a servicios especificos.",
    )
    add_table(
        doc,
        ["Componente", "Archivo o servicio", "Funcion"],
        [
            ["Estructura del sitio", "index.html, nuestra_empresa.html, beneficio.html, trabaja.html, contacto.html", "Divide el contenido por secciones claras para navegacion publica."],
            ["Diseno visual", "style.css", "Define identidad visual, layout, hero, tarjetas, formularios, responsive y adaptacion movil."],
            ["Chatbot", "js/chatbot.js y css/chatbot.css", "Muestra asistente flotante y redirige a WhatsApp segun la necesidad del visitante."],
            ["Contacto", "FormSubmit", "Envia mensajes del formulario de contacto al correo configurado."],
            ["Postulaciones", "EmailJS y Cloudinary", "Envia datos del candidato y guarda hoja de vida cargada por el usuario."],
            ["Ubicacion", "Google Maps iframe", "Permite ver la ubicacion de la empresa dentro de la pagina."],
            ["Publicacion", "ZIP para Hostinger", "Empaqueta el sitio para subirlo al hosting."],
        ],
        header_fill=ACCENT_2,
    )

    add_heading(doc, "24.4 Por que se hizo de esa manera", 2)
    add_table(
        doc,
        ["Decision", "Razon"],
        [
            ["Sitio estatico con HTML, CSS y JavaScript", "La web institucional no necesita una base de datos propia para mostrar contenido; esto simplifica mantenimiento y publicacion."],
            ["Paginas separadas por tema", "Hace mas facil que un visitante encuentre informacion sobre empresa, beneficio, vacantes o contacto."],
            ["Diseno responsive", "El sitio debe verse bien en computador y celular, porque muchos usuarios llegaran desde WhatsApp o redes sociales."],
            ["Chatbot a WhatsApp", "WhatsApp es un canal directo y familiar para visitantes, clientes y candidatos."],
            ["Formulario laboral con servicios externos", "Permite recibir postulaciones sin crear todavia un backend propio para recursos humanos."],
            ["Carga de hoja de vida a Cloudinary", "Soluciona el envio de archivos y permite incluir un enlace de descarga en el correo recibido."],
            ["FormSubmit en contacto", "Permite habilitar mensajes de contacto de forma rapida y funcional."],
            ["ZIP para Hostinger", "Facilita subir el proyecto al hosting como sitio publico."],
        ],
        header_fill=ACCENT_2,
    )

    add_heading(doc, "24.5 Flujos funcionales implementados", 2)
    add_table(
        doc,
        ["Flujo", "Como lo usa una persona", "Resultado"],
        [
            ["Navegacion institucional", "El visitante entra al inicio y se mueve por Nuestra Empresa, Beneficio, Trabaja y Contacto.", "Conoce la empresa, servicios, historia y canales de atencion."],
            ["Contacto", "El visitante llena nombre, correo y mensaje.", "El mensaje se envia al correo configurado con FormSubmit."],
            ["Postulacion laboral", "El candidato abre el modal, completa datos, acepta tratamiento de datos y adjunta hoja de vida.", "El sistema calcula puntaje, clasifica al candidato, sube el archivo y envia la postulacion."],
            ["Atencion por WhatsApp", "El usuario abre el chatbot y elige hojas de vida o servicios.", "La web abre WhatsApp con un mensaje prellenado al numero correspondiente."],
            ["Ubicacion", "El usuario consulta el mapa embebido.", "Puede ubicar la empresa en Armenia, Quindio."],
        ],
        header_fill=ACCENT_2,
    )

    add_heading(doc, "24.6 Evidencia tecnica del trabajo realizado", 2)
    add_table(
        doc,
        ["Elemento", "Evidencia en la carpeta del proyecto"],
        [
            ["Paginas HTML", "index.html, nuestra_empresa.html, beneficio.html, trabaja.html y contacto.html."],
            ["Estilos", "style.css y css/chatbot.css con reglas responsive, menu movil, modales, tarjetas y chatbot."],
            ["Interactividad", "js/chatbot.js y scripts en trabaja.html para modal, campos condicionales, puntaje y envio."],
            ["Recursos visuales", "Carpeta img con logotipo, fotos institucionales, fondos y recursos graficos."],
            ["Publicacion", "Archivo agro-angeles-hostinger-actualizado-20260826.zip preparado para hosting."],
        ],
        header_fill=ACCENT_2,
    )

    add_heading(doc, "24.7 Evidencia visual", 2)
    add_image(doc, WEB_IMG / "Logo 1.png", "Identidad grafica usada en la pagina web", width=1.55)
    add_image(doc, collage, "Resumen visual de activos y secciones de la pagina web", width=6.35)

    add_heading(doc, "24.8 Resultado", 2)
    add_paragraph(
        doc,
        "El resultado es una pagina web institucional lista para presentar la empresa, atender visitantes y recibir postulaciones. "
        "El desarrollo muestra trabajo en estructura, diseno responsive, experiencia de usuario, integracion con servicios externos y preparacion para publicacion en hosting.",
    )
    add_paragraph(
        doc,
        "Este proyecto demuestra una linea distinta de trabajo frente a los sistemas internos: mientras el frontend modular y el cronograma resuelven procesos operativos, "
        "la pagina web fortalece la presencia publica de Agro Angeles y crea canales digitales de contacto y reclutamiento.",
    )


def main():
    doc = Document(INPUT)
    insert_toc_entries(doc)
    append_part_divider(doc)
    append_cronograma_section(doc)
    append_web_section(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
